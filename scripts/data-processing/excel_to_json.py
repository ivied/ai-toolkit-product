#!/usr/bin/env python3
"""
Excel/CSV to JSON Converter — Convert spreadsheet data to clean JSON.

Handles CSV and basic Excel-like formats, with options for nested output,
type inference, and field mapping.

Usage:
    python excel_to_json.py input.csv
    python excel_to_json.py input.csv --output data.json --pretty
    python excel_to_json.py input.csv --nest "address=street,city,zip"
    python excel_to_json.py input.csv --rename "old_name=new_name"

No external dependencies for CSV. Install openpyxl for .xlsx support.
"""

import argparse
import csv
import json
import re
import sys
from pathlib import Path


def infer_type(value):
    if not value or value.strip() == "": return None
    v = value.strip()
    if v.lower() in ("true", "false"): return v.lower() == "true"
    try: return int(v)
    except ValueError: pass
    try: return float(v)
    except ValueError: pass
    return v


def read_csv(filepath, delimiter=","):
    with open(filepath, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        return list(reader), list(reader.fieldnames or [])


def read_xlsx(filepath, sheet=0):
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("Install openpyxl for .xlsx: pip install openpyxl", file=sys.stderr)
        sys.exit(1)
    
    wb = load_workbook(filepath, read_only=True)
    ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
    
    rows = list(ws.iter_rows(values_only=True))
    if not rows: return [], []
    
    headers = [str(h or f"col_{i}") for i, h in enumerate(rows[0])]
    data = [dict(zip(headers, row)) for row in rows[1:]]
    return data, headers


def apply_nesting(row, nest_rules):
    """Nest flat fields into objects. Rule: parent=child1,child2"""
    for rule in nest_rules:
        parent, fields_str = rule.split("=", 1)
        fields = [f.strip() for f in fields_str.split(",")]
        nested = {}
        for f in fields:
            if f in row:
                nested[f] = row.pop(f)
        if nested:
            row[parent] = nested
    return row


def apply_rename(row, rename_rules):
    """Rename fields. Rule: old=new"""
    for rule in rename_rules:
        old, new = rule.split("=", 1)
        if old in row:
            row[new] = row.pop(old)
    return row


def convert(filepath, type_infer=True, nest=None, rename=None, skip_empty=True):
    ext = Path(filepath).suffix.lower()
    
    if ext in (".xlsx", ".xls"):
        data, headers = read_xlsx(filepath)
    elif ext == ".tsv":
        data, headers = read_csv(filepath, delimiter="\t")
    else:
        data, headers = read_csv(filepath)
    
    result = []
    for row in data:
        if skip_empty and all(not v for v in row.values()):
            continue
        
        processed = {}
        for k, v in row.items():
            if type_infer:
                processed[k] = infer_type(str(v) if v is not None else "")
            else:
                processed[k] = v
        
        if rename:
            processed = apply_rename(processed, rename)
        if nest:
            processed = apply_nesting(processed, nest)
        
        result.append(processed)
    
    return result, headers


def main():
    parser = argparse.ArgumentParser(description="Convert CSV/Excel to JSON")
    parser.add_argument("input", help="Input file (CSV, TSV, XLSX)")
    parser.add_argument("--output", "-o", help="Output JSON file")
    parser.add_argument("--pretty", action="store_true", default=True)
    parser.add_argument("--compact", action="store_true")
    parser.add_argument("--no-type-infer", action="store_true")
    parser.add_argument("--nest", action="append", help="Nest fields: parent=child1,child2")
    parser.add_argument("--rename", action="append", help="Rename field: old=new")
    parser.add_argument("--skip-empty", action="store_true", default=True)
    parser.add_argument("--wrap", help="Wrap array in object with this key")
    args = parser.parse_args()
    
    data, headers = convert(args.input, not args.no_type_infer, args.nest, args.rename, args.skip_empty)
    
    output = data
    if args.wrap:
        output = {args.wrap: data, "count": len(data)}
    
    indent = None if args.compact else 2
    json_str = json.dumps(output, indent=indent, ensure_ascii=False, default=str)
    
    if args.output:
        with open(args.output, "w") as f:
            f.write(json_str)
        print(f"✅ Converted {len(data)} rows → {args.output}", file=sys.stderr)
    else:
        print(json_str)


if __name__ == "__main__":
    main()
